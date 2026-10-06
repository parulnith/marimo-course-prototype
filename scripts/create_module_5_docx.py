from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "course/modules/05-reusable-systems-studio-draft.mdx"
OUTPUT = ROOT / "output/module-5-from-prototype-to-production.docx"

IMAGES = {
    "studioNotebookViews": ROOT / "course/images/module-5/studio-notebook-views.webp",
    "studioAddSimpleView": ROOT / "course/images/module-5/studio-add-simple-view.png",
    "studioSimpleView": ROOT / "course/images/module-5/studio-simple-view.png",
    "occupancyReport": ROOT / "course/images/module-5/occupancy-report.png",
    "appViewGif": ROOT / "course/images/module-5/app-view.gif",
    "exportFormatsGif": ROOT / "course/images/module-5/export-formats.gif",
}


def shade(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_margins(paragraph, top: int = 90, bottom: int = 90, left: int = 120, right: int = 120) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), str(top))
    spacing.set(qn("w:after"), str(bottom))


def add_inline(paragraph, text: str) -> None:
    text = re.sub(r'<a href="([^"]+)">([^<]+)</a>', r'\2 (\1)', text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"<code>(.*?)</code>", r"`\1`", text)
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
        else:
            run.text = part


def add_body(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    add_inline(paragraph, text.strip())


def add_code(doc: Document, code: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    shade(paragraph, "F3F4F6")
    set_cell_margins(paragraph)
    run = paragraph.add_run(code)
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    run.font.size = Pt(8.5)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].paragraph_format.space_before = Pt(16)
        styles[name].paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph(style="Title")
    title.add_run("Module 5 From Prototype to Production")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Course draft for Google Docs").italic = True

    lines = SOURCE.read_text().splitlines()
    in_front_matter = False
    in_code = False
    code_lines: list[str] = []
    in_callout = False
    in_try_it = False
    in_embed = False
    in_video = False
    in_quiz = False
    quiz_lines: list[str] = []

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if stripped == "---":
            in_front_matter = not in_front_matter
            continue
        if in_front_matter or stripped.startswith("import "):
            continue

        if stripped.startswith("```"):
            if in_code:
                add_code(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(raw)
            continue

        if stripped.startswith("<MarimoEmbed"):
            in_embed = True
            continue
        if in_embed:
            match = re.search(r'title="([^"]+)"', stripped)
            if match:
                paragraph = doc.add_paragraph()
                paragraph.add_run("Interactive notebook: ").bold = True
                paragraph.add_run(match.group(1))
            if stripped.endswith("/>"):
                in_embed = False
            continue

        if stripped.startswith("<Callout"):
            title_match = re.search(r'title="([^"]+)"', stripped)
            paragraph = doc.add_paragraph()
            paragraph.add_run(title_match.group(1) if title_match else "Note").bold = True
            shade(paragraph, "EAF2F8")
            in_callout = True
            continue
        if stripped == "</Callout>":
            in_callout = False
            continue
        if stripped == "<TryIt>":
            paragraph = doc.add_paragraph()
            paragraph.add_run("Try it").bold = True
            shade(paragraph, "FDF2E9")
            in_try_it = True
            continue
        if stripped == "</TryIt>":
            in_try_it = False
            continue

        if stripped.startswith("<video"):
            in_video = True
            continue
        if in_video:
            source = re.search(r'src=\{([^}]+)\}', stripped)
            if source:
                paragraph = doc.add_paragraph()
                paragraph.add_run("Video: ").bold = True
                paragraph.add_run(source.group(1).replace("Video", " video"))
                paragraph.runs[-1].italic = True
            if stripped == "</video>":
                in_video = False
            continue

        if stripped.startswith("<Quiz"):
            in_quiz = True
            quiz_lines = [stripped]
            continue
        if in_quiz:
            quiz_lines.append(stripped)
            if stripped.endswith("/>"):
                blob = " ".join(quiz_lines)
                question = re.search(r'question="([^"]+)"', blob)
                options = re.findall(r'"([^"]+)"', re.search(r'options=\{\[(.*?)\]\}', blob).group(1))
                answer = int(re.search(r'answer=\{(\d+)\}', blob).group(1))
                paragraph = doc.add_paragraph()
                paragraph.add_run("Check your understanding: ").bold = True
                paragraph.add_run(question.group(1))
                for index, option in enumerate(options):
                    p = doc.add_paragraph(style="List Bullet")
                    p.add_run(option)
                    if index == answer:
                        p.add_run(" (correct answer)").italic = True
                in_quiz = False
            continue

        image_match = re.search(r'src=\{([^}]+)\}', stripped)
        if image_match:
            image = IMAGES.get(image_match.group(1))
            if image and image.exists():
                if image.suffix.lower() == ".webp":
                    converted = ROOT / "output" / f"{image.stem}.png"
                    converted.parent.mkdir(parents=True, exist_ok=True)
                    Image.open(image).convert("RGB").save(converted)
                    image = converted
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image), width=Inches(6.4))
            continue

        if stripped.startswith("{/*") and stripped.endswith("*/}"):
            note = stripped[3:-3].strip()
            if note:
                paragraph = doc.add_paragraph()
                paragraph.add_run(note).italic = True
                shade(paragraph, "FFF2CC")
            continue
        if stripped.startswith("<p") or stripped == "</p>":
            text = re.sub(r"<[^>]+>", "", stripped)
            if text:
                add_body(doc, text)
            continue
        if not stripped:
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            style = {1: "Heading 1", 2: "Heading 1", 3: "Heading 2"}[level]
            add_body(doc, heading.group(2), style)
            continue
        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            continue
        if in_callout or in_try_it:
            paragraph = doc.add_paragraph()
            shade(paragraph, "EAF2F8" if in_callout else "FDF2E9")
            add_inline(paragraph, stripped)
            continue
        add_body(doc, stripped)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Module 5 From Prototype to Production"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
